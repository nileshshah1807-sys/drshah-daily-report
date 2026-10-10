"""
Export builders — branded PDF reports, indicators CSV and formatted Excel file.
"""
import csv
import io
from datetime import datetime
from xml.sax.saxutils import escape

import analyzer
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import (HRFlowable, Image, KeepTogether, Paragraph,
                                SimpleDocTemplate, Spacer, Table, TableStyle)

import paths
BASE = paths.BUNDLE_DIR
NAVY = colors.HexColor("#0b1220")
GOLD = colors.HexColor("#f5b52e")
RATING_COLORS = {
    "Strong Buy": colors.HexColor("#059669"),
    "Buy": colors.HexColor("#22c55e"),
    "Hold": colors.HexColor("#f59e0b"),
    "Sell": colors.HexColor("#f97316"),
    "Strong Sell": colors.HexColor("#ef4444"),
}
RATING_HEX = {k: "#" + c.hexval()[2:] for k, c in RATING_COLORS.items()}
LOGO = paths.LOGO_PATH


def _pdf_styles():
    ss = getSampleStyleSheet()
    return {
        "h1": ParagraphStyle("h1", parent=ss["Title"], fontName="Helvetica-Bold",
                             fontSize=20, textColor=NAVY, leading=24),
        "h2": ParagraphStyle("h2", parent=ss["Heading2"], fontName="Helvetica-Bold",
                             fontSize=13, textColor=NAVY, spaceBefore=6, spaceAfter=4),
        "body": ParagraphStyle("body", parent=ss["BodyText"], fontName="Helvetica",
                               fontSize=9, textColor=colors.HexColor("#334155"), leading=12),
        "small": ParagraphStyle("small", parent=ss["BodyText"], fontName="Helvetica",
                                fontSize=8, textColor=colors.HexColor("#64748b"), leading=10),
    }


def _brand_header(doc, title, subtitle, text_width=160 * mm):
    logo = Image(LOGO, width=14 * mm, height=14 * mm)
    logo.hAlign = "LEFT"
    tbl = Table([[logo, [Paragraph(title, doc["h1"]),
                          Paragraph(subtitle, doc["small"])]]],
                colWidths=[18 * mm, text_width])
    tbl.setStyle(TableStyle([
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 0),
    ]))
    return tbl


def _fmt(v, suffix=""):
    if not isinstance(v, (int, float)):
        return "—"
    text = f"{v:,.2f}"
    return ("0.00" if text == "-0.00" else text) + suffix      # never "-0.00"


def _vol(v) -> str:
    """Compact share volume: 134914600 → '134.9M' (the full number of a busy
    stock is too wide for its table cell)."""
    if not isinstance(v, (int, float)) or v != v:
        return "—"
    for limit, suffix in ((1e9, "B"), (1e6, "M"), (1e3, "K")):
        if abs(v) >= limit:
            return f"{v / limit:.1f}{suffix}"
    return f"{v:.0f}"


def _by_rating_then_score(results: dict) -> list:
    """Best first: rating, then the score inside each rating."""
    return sorted(results.items(),
                  key=lambda kv: (analyzer.rating_rank(kv[1]["rating"]), kv[1]["score"]),
                  reverse=True)


def _pts(v) -> str:
    return f"{v:+.1f} pts" if isinstance(v, (int, float)) else "—"


def _fund_label(fund: dict) -> str | None:
    if not fund:
        return None
    if not fund.get("applicable", True):
        return f"n/a ({str(fund.get('type', 'fund')).title()})"
    return f"{fund.get('label')} ({fund.get('passed')}/{fund.get('total')})"


def _fundamentals_line(fund: dict) -> str:
    if not fund:
        return ""
    if not fund.get("applicable", True):
        dy = fund.get("dividend_yield")
        return ("<b>Fundamentals:</b> not applicable (" + str(fund.get("type", "fund")).title() + ")"
                + (f" · yield {dy:.2f}%" if isinstance(dy, (int, float)) else ""))
    bits = []
    for key, label, fmt in (("pe", "P/E", "{:.1f}"), ("revenue_growth", "revenue", "{:+.1f}%"),
                            ("profit_margin", "margin", "{:.1f}%"), ("roe", "ROE", "{:.1f}%"),
                            ("debt_to_equity", "debt/equity", "{:.2f}"),
                            ("dividend_yield", "dividend", "{:.2f}%")):
        v = fund.get(key)
        if isinstance(v, (int, float)):
            bits.append(f"{label} {fmt.format(v)}")
    return ("<b>Fundamentals:</b> " + " · ".join(bits)
            + f" — {fund.get('label')} ({fund.get('passed')}/{fund.get('total')} checks)")


def _extras_lines(r: dict) -> list[str]:
    """Fundamentals, relative strength, score trend and earnings for one
    stock in the report PDF — one short line each."""
    parts = []
    fl = _fundamentals_line(r.get("fundamentals") or {})
    if fl:
        parts.append(fl)
    rs = r.get("relative_strength") or {}
    if rs.get("rs_6m") is not None:
        adj = rs.get("score_adj") or 0
        parts.append(f"<b>vs S&amp;P 500:</b> 6M {_pts(rs.get('rs_6m'))} · 12M {_pts(rs.get('rs_12m'))}"
                     f" — {rs.get('label', '')}" + (f" ({adj:+.0f} score)" if adj else ""))
    tr = r.get("score_trend") or {}
    ago = [f"{name} {tr[k]:.1f}" for k, name in (("w1", "1 week ago"), ("m1", "1 month ago"),
                                                 ("m3", "3 months ago"))
           if isinstance(tr.get(k), (int, float))]
    if ago:
        parts.append("<b>Score trend:</b> " + " · ".join(ago))
    earn = r.get("earnings") or {}
    if earn.get("date"):
        txt = f"<b>Next earnings:</b> {earn['date']}"
        if earn.get("soon"):
            txt = f"<font color='#dc2626'>{txt} — results due within a week</font>"
        parts.append(txt)
    return parts


def _rating_summary_line(rating_counts: dict, style) -> Paragraph:
    return Paragraph(
        "Summary: " + "  |  ".join(
            f"{k}: <font color='{RATING_HEX[k]}'><b>{v}</b></font>"
            for k, v in rating_counts.items()), style)


# ---------------------------------------------------------------------------
# Watchlist PDF
# ---------------------------------------------------------------------------

def build_watchlist_pdf(results: dict, errors: dict) -> bytes:
    buf = io.BytesIO()
    # LANDSCAPE: the 12-column table is 257 mm wide — on a portrait page
    # (178 mm usable) its first and last columns were cut off at the edges
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4),
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Watchlist Analysis — Dr. Shah's US Stocks Analysis")
    st = _pdf_styles()
    story = []
    story.append(_brand_header(
        st, "Dr. Shah's US Stocks Analysis",
        f"Technical Watchlist Report  •  Generated {datetime.now().strftime('%d %b %Y, %I:%M %p')}",
        text_width=245 * mm))
    story.append(HRFlowable(width="100%", thickness=1.2, color=GOLD, spaceAfter=8))

    rating_counts = {k: 0 for k in RATING_COLORS}
    for r in results.values():
        rating_counts[r["rating"]] = rating_counts.get(r["rating"], 0) + 1
    story.append(_rating_summary_line(rating_counts, st["body"]))
    story.append(Spacer(1, 6))

    rows = [["#", "Ticker", "Company", "Price ($)", "Change %", "Score", "Rating",
             "Long-Term Trend", "Momentum", "RSI", "vs S&P 6M", "Daily Pattern"]]
    for i, (tk, r) in enumerate(_by_rating_then_score(results), 1):
        sig = {s["name"]: s for s in r["signals"]}
        rs6 = (r.get("relative_strength") or {}).get("rs_6m")
        rows.append([
            str(i), tk, r["name"][:32], _fmt(r["price"]),
            f"{r['change_pct']:+.2f}", f"{r['score']:.1f}", r["rating"],
            sig["Long-Term Trend"]["label"], sig["Investment Momentum"]["label"],
            _fmt(r["indicators"].get("rsi")),
            f"{rs6:+.1f}" if isinstance(rs6, (int, float)) else "—", r["pattern"][:32],
        ])
    tbl = Table(rows, colWidths=[8 * mm, 16 * mm, 50 * mm, 18 * mm, 16 * mm, 13 * mm,
                                 21 * mm, 28 * mm, 30 * mm, 12 * mm, 17 * mm, 28 * mm],
                repeatRows=1)
    style = [
        ("BACKGROUND", (0, 0), (-1, 0), NAVY),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("FONTSIZE", (0, 0), (-1, -1), 7.5),
        ("FONTNAME", (1, 1), (1, -1), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#cbd5e1")),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f1f5f9")]),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 3),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]
    for ri in range(1, len(rows)):
        rating = rows[ri][6]
        if rating in RATING_COLORS:
            style.append(("BACKGROUND", (6, ri), (6, ri), RATING_COLORS[rating]))
            style.append(("TEXTCOLOR", (6, ri), (6, ri), colors.white))
            style.append(("FONTNAME", (6, ri), (6, ri), "Helvetica-Bold"))
    tbl.setStyle(TableStyle(style))
    story.append(tbl)

    if errors:
        story.append(Spacer(1, 8))
        story.append(Paragraph("<b>Note:</b> symbols with no data (skipped): "
                               + escape(", ".join(f"{k} ({v})" for k, v in errors.items())),
                               st["body"]))
    story.append(Spacer(1, 8))
    story.append(Paragraph(
        "Rating scale — Strong Buy ≥ 80 • Buy 65–79 • Hold 50–64 • Sell 35–49 • Strong Sell < 35.  "
        "Analysis is technical only and not investment advice.", st["small"]))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Full report PDF
# ---------------------------------------------------------------------------

def build_report_pdf(results: dict, errors: dict) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4,
                            leftMargin=16 * mm, rightMargin=16 * mm,
                            topMargin=14 * mm, bottomMargin=14 * mm,
                            title="Full Analysis Report — Dr. Shah's US Stocks Analysis")
    st = _pdf_styles()
    story = []
    story.append(_brand_header(
        st, "Dr. Shah's US Stocks Analysis",
        f"Full Technical Analysis Report  •  {datetime.now().strftime('%d %b %Y, %I:%M %p')}"))
    story.append(HRFlowable(width="100%", thickness=1.2, color=GOLD, spaceAfter=8))

    rating_counts = {k: 0 for k in RATING_COLORS}
    avg = 0.0
    for r in results.values():
        rating_counts[r["rating"]] += 1
        avg += r["score"]
    n = len(results) or 1
    avg /= n
    story.append(Paragraph(f"<b>Coverage:</b> {len(results)} stocks  •  "
                           f"<b>Average score:</b> {avg:.1f}/100", st["body"]))
    story.append(Spacer(1, 4))
    story.append(_rating_summary_line(rating_counts, st["body"]))
    story.append(Spacer(1, 8))

    for tk, r in _by_rating_then_score(results):
        sig = {s["name"]: s for s in r["signals"]}
        ind = r["indicators"]
        # heading + summary + table stay on one page (a heading used to be
        # stranded at the foot of a page with its table on the next)
        block = [Paragraph(
            f"{escape(tk)} — {escape(str(r['name']))}  "
            f"<font color='{RATING_HEX[r['rating']]}'>[{r['rating']}]</font>", st["h2"])]
        block.append(Paragraph(
            f"Price <b>${r['price']:,.2f}</b> ({r['change_pct']:+.2f}% today)  •  "
            f"Score <b>{r['score']:.1f}/100</b>  •  Market cap {r['market_cap']}  •  "
            f"Last daily candle: {escape(str(r['pattern']))}", st["body"]))
        # (no emoji: the PDF font has none and printed a black square instead)
        block.append(Paragraph(f"<b>Verdict:</b> {escape(str(r.get('verdict', '')))}", st["body"]))
        # what a verdict is (and is not) belongs next to it, not only in the footer
        notes = [str(r[k]) for k in ("disclaimer", "rating_note") if r.get(k)]
        if notes:
            block.append(Paragraph(escape(" ".join(notes)), st["small"]))
        for line in _extras_lines(r):
            block.append(Paragraph(line, st["body"]))
        block.append(Spacer(1, 4))

        sr = r["support_resistance"]
        ind_rows = [
            ["Metric", "Value", "Signal", "Metric", "Value", "Signal"],
            ["RSI (14)", _fmt(ind.get("rsi")), sig["Investment Momentum"]["label"],
             "SMA 50 / 150 / 260", f"{_fmt(ind.get('sma50'))} / {_fmt(ind.get('sma150'))} / {_fmt(ind.get('sma260'))}",
             sig["Long-Term Trend"]["label"]],
            ["MACD line", _fmt(ind.get("macd_line")), sig["Investment Momentum"]["label"],
             "3M / 6M return %", f"{_fmt(ind.get('roc63'))} / {_fmt(ind.get('roc126'))}", "—"],
            ["MACD histogram", _fmt(ind.get("macd_hist")), "—",
             "ATR (14) %", _fmt(ind.get("atr_pct")), sig["Stability"]["label"]],
            ["SMA 200 (daily)", _fmt(ind.get("sma200")), sig["Golden Cross"]["label"],
             "Bollinger (20, 2σ)", f"{_fmt(ind.get('bb_lower'))} – {_fmt(ind.get('bb_upper'))}", "—"],
            ["Support (nearest)", _fmt(sr["nearest_support"]),
             f"{(sr['dist_to_support_pct'] or 0):.1f}% away",
             "Resistance (nearest)", _fmt(sr["nearest_resistance"]),
             f"{(sr['dist_to_resistance_pct'] or 0):.1f}% away"],
            ["52W high / low", f"{_fmt(r['week52']['high'])} / {_fmt(r['week52']['low'])}",
             sig["52W Position"]["label"],
             "Volume (today / 50d avg)",
             f"{_vol(ind.get('vol_today'))} / {_vol(ind.get('vol_avg50'))}",
             sig["Volume"]["label"]],
        ]
        # label columns wide enough for their text: "Resistance (nearest)" and
        # "Volume (today / 50d avg)" used to run into the value next to them
        it = Table(ind_rows, colWidths=[28 * mm, 24 * mm, 31 * mm, 36 * mm, 33 * mm, 26 * mm])
        it.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), NAVY),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("FONTSIZE", (0, 0), (-1, -1), 7.5),
            ("GRID", (0, 0), (-1, -1), 0.3, colors.HexColor("#cbd5e1")),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
            ("TOPPADDING", (0, 0), (-1, -1), 2.5),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 2.5),
        ]))
        block.append(it)
        block.append(Spacer(1, 6))
        story.append(KeepTogether(block))
        for s in r["signals"]:
            story.append(Paragraph(
                f"<b>{s['name']} ({s['score']:.0f}/100, weight {s['weight']*100:.0f}%):</b> "
                f"{s['label']} — {escape(str(s['detail']))}", st["small"]))
        story.append(Spacer(1, 8))

    if errors:
        story.append(Paragraph("<b>Note:</b> symbols with no data: " + escape(", ".join(errors)),
                               st["small"]))
    story.append(HRFlowable(width="100%", thickness=0.8, color=GOLD, spaceBefore=4, spaceAfter=4))
    story.append(Paragraph(
        "Generated by Dr. Shah's US Stocks Analysis (3-year daily charts). Indicators: RSI(14), "
        "MACD(12,26,9), Bollinger(20,2σ), ATR(14), SMA(20/50/150/200/260), 3M/6M returns, "
        "volume accumulation, 52-week position, daily candlestick patterns. "
        "This report is for educational purposes only and is not investment advice.", st["small"]))
    doc.build(story)
    return buf.getvalue()


# ---------------------------------------------------------------------------
# Indicators CSV
# ---------------------------------------------------------------------------

INDICATOR_COLUMNS = [
    ("ticker", "Ticker"), ("name", "Company"), ("price", "Price ($)"),
    ("change_pct", "Change %"), ("score", "Score"), ("rating", "Rating"),
    ("rsi", "RSI(14)"), ("macd_line", "MACD Line"), ("macd_signal", "MACD Signal"),
    ("macd_hist", "MACD Histogram"), ("sma20", "SMA(20)"), ("sma50", "SMA(50)"),
    ("sma150", "SMA(150)"), ("sma200", "SMA(200)"), ("sma260", "SMA(260)"),
    ("atr_pct", "ATR(14) %"), ("roc63", "3-Month Return %"), ("roc126", "6-Month Return %"),
    ("bb_upper", "Bollinger Upper"), ("bb_mid", "Bollinger Mid"), ("bb_lower", "Bollinger Lower"),
    ("stoch_k", "Stochastic %K"), ("stoch_d", "Stochastic %D"),
    ("vol_today", "Volume Today"), ("vol_avg50", "Volume Avg(50d)"), ("vol_avg200", "Volume Avg(200d)"),
    ("lt_score", "Long-Term Trend Score"), ("lt_label", "Long-Term Trend Label"),
    ("mom_score", "Momentum Score"), ("mom_label", "Momentum Label"),
    ("gc_score", "Golden Cross Score"), ("gc_label", "Golden Cross Label"),
    ("st_score", "Stability Score"), ("st_label", "Stability Label"),
    ("vol_score", "Volume Score"), ("vol_label", "Volume Label"),
    ("pos_score", "52W Position Score"), ("pos_label", "52W Position Label"),
    ("pattern", "Daily Pattern"), ("verdict", "Investment Verdict"),
    ("support", "Nearest Support"), ("resistance", "Nearest Resistance"),
    ("w52_high", "52W High"), ("w52_low", "52W Low"),
    ("rs_6m", "vs S&P 500 6M (pts)"), ("rs_12m", "vs S&P 500 12M (pts)"),
    ("rs_label", "Relative Strength"),
    ("score_1w", "Score 1 Week Ago"), ("score_1m", "Score 1 Month Ago"),
    ("next_earnings", "Next Earnings"),
    ("pe", "P/E"), ("forward_pe", "Forward P/E"),
    ("revenue_growth", "Revenue Growth %"), ("earnings_growth", "Earnings Growth %"),
    ("profit_margin", "Profit Margin %"), ("roe", "Return on Equity %"),
    ("debt_to_equity", "Debt/Equity"), ("dividend_yield", "Dividend Yield %"),
    ("fund_label", "Fundamentals"),
    ("last_updated", "Last Updated"),
]


def _flatten_result(r: dict) -> dict:
    sig = {s["name"]: s for s in r["signals"]}
    ind = r["indicators"]
    sr = r["support_resistance"]
    rs = r.get("relative_strength") or {}
    tr = r.get("score_trend") or {}
    earn = r.get("earnings") or {}
    fund = r.get("fundamentals") or {}
    return {
        "ticker": r["ticker"], "name": r["name"], "price": r["price"],
        "change_pct": r["change_pct"], "score": r["score"], "rating": r["rating"],
        "rsi": ind.get("rsi"), "macd_line": ind.get("macd_line"),
        "macd_signal": ind.get("macd_signal"), "macd_hist": ind.get("macd_hist"),
        "sma20": ind.get("sma20"), "sma50": ind.get("sma50"),
        "sma150": ind.get("sma150"), "sma200": ind.get("sma200"),
        "sma260": ind.get("sma260"), "atr_pct": ind.get("atr_pct"),
        "roc63": ind.get("roc63"), "roc126": ind.get("roc126"),
        "bb_upper": ind.get("bb_upper"), "bb_mid": ind.get("bb_mid"),
        "bb_lower": ind.get("bb_lower"), "stoch_k": ind.get("stoch_k"),
        "stoch_d": ind.get("stoch_d"),
        "vol_today": ind.get("vol_today"), "vol_avg50": ind.get("vol_avg50"),
        "vol_avg200": ind.get("vol_avg200"),
        "lt_score": sig["Long-Term Trend"]["score"], "lt_label": sig["Long-Term Trend"]["label"],
        "mom_score": sig["Investment Momentum"]["score"], "mom_label": sig["Investment Momentum"]["label"],
        "gc_score": sig["Golden Cross"]["score"], "gc_label": sig["Golden Cross"]["label"],
        "st_score": sig["Stability"]["score"], "st_label": sig["Stability"]["label"],
        "vol_score": sig["Volume"]["score"], "vol_label": sig["Volume"]["label"],
        "pos_score": sig["52W Position"]["score"], "pos_label": sig["52W Position"]["label"],
        "pattern": r["pattern"], "verdict": r.get("verdict", ""),
        "support": sr.get("nearest_support"), "resistance": sr.get("nearest_resistance"),
        "w52_high": r["week52"]["high"], "w52_low": r["week52"]["low"],
        "rs_6m": rs.get("rs_6m"), "rs_12m": rs.get("rs_12m"), "rs_label": rs.get("label"),
        "score_1w": tr.get("w1"), "score_1m": tr.get("m1"),
        "next_earnings": earn.get("date"),
        "pe": fund.get("pe"), "forward_pe": fund.get("forward_pe"),
        "revenue_growth": fund.get("revenue_growth"),
        "earnings_growth": fund.get("earnings_growth"),
        "profit_margin": fund.get("profit_margin"), "roe": fund.get("roe"),
        "debt_to_equity": fund.get("debt_to_equity"),
        "dividend_yield": fund.get("dividend_yield"),
        "fund_label": _fund_label(fund),
        "last_updated": r["last_updated"],
    }


def build_indicators_csv(results: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow([label for _, label in INDICATOR_COLUMNS])
    for _tk, r in sorted(results.items(),
                         key=lambda kv: analyzer.rating_rank(kv[1]["rating"]), reverse=True):
        row = _flatten_result(r)
        w.writerow([row[key] for key, _ in INDICATOR_COLUMNS])
    return buf.getvalue().encode("utf-8-sig")


# ---------------------------------------------------------------------------
# Indicators Excel (.xlsx)
# ---------------------------------------------------------------------------

def build_indicators_xlsx(results: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Indicators"

    header_fill = PatternFill("solid", fgColor="0B1220")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    rating_fills = {
        "Strong Buy": PatternFill("solid", fgColor="059669"),
        "Buy": PatternFill("solid", fgColor="22C55E"),
        "Hold": PatternFill("solid", fgColor="F59E0B"),
        "Sell": PatternFill("solid", fgColor="F97316"),
        "Strong Sell": PatternFill("solid", fgColor="EF4444"),
    }

    for ci, (_key, label) in enumerate(INDICATOR_COLUMNS, start=1):
        cell = ws.cell(row=1, column=ci, value=label)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for ri, (_tk, r) in enumerate(sorted(results.items(),
                                         key=lambda kv: analyzer.rating_rank(kv[1]["rating"]), reverse=True), start=2):
        row = _flatten_result(r)
        for ci, (key, _label) in enumerate(INDICATOR_COLUMNS, start=1):
            cell = ws.cell(row=ri, column=ci, value=row[key])
            if key in ("price", "score", "change_pct", "rsi", "roc63", "roc126",
                       "atr_pct", "stoch_k", "stoch_d", "lt_score", "mom_score",
                       "gc_score", "st_score", "vol_score", "pos_score",
                       "rs_6m", "rs_12m", "score_1w", "score_1m", "pe", "forward_pe",
                       "revenue_growth", "earnings_growth", "profit_margin", "roe",
                       "debt_to_equity", "dividend_yield"):
                cell.number_format = "0.00"
            if key == "rating":
                cell.fill = rating_fills.get(row["rating"], PatternFill())
                cell.font = Font(bold=True, color="FFFFFF")
        ws.cell(row=ri, column=6).fill = rating_fills.get(row["rating"], PatternFill())

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(INDICATOR_COLUMNS))}{len(results) + 1}"
    for ci in range(1, len(INDICATOR_COLUMNS) + 1):
        col = get_column_letter(ci)
        width = max(len(INDICATOR_COLUMNS[ci - 1][1]) + 2, 10)
        for cell in ws[col][1:len(results) + 2]:
            v = cell.value
            if isinstance(v, (int, float)):
                width = max(width, len(f"{v:,.2f}") + 2)
            elif v is not None:
                width = max(width, min(len(str(v)) + 2, 26))
        ws.column_dimensions[col].width = min(width, 30)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
